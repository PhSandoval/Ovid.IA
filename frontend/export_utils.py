import io
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def gerar_docx_parecer(pergunta, parecer_texto):
    """Gera um arquivo .docx formatado para um parecer de RAG"""
    doc = Document()
    
    # Título principal
    titulo = doc.add_heading('Parecer Jurídico - Ovid.IA', 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Pergunta
    doc.add_heading('1. Tese / Pergunta Formulada', level=1)
    p_perg = doc.add_paragraph(pergunta)
    p_perg.style.font.name = 'Arial'
    
    # Resposta
    doc.add_heading('2. Análise Jurisprudencial', level=1)
    p_resp = doc.add_paragraph(parecer_texto)
    p_resp.style.font.name = 'Arial'
    
    # Rodapé
    doc.add_paragraph('\n')
    rodape = doc.add_paragraph('Gerado por Ovid.IA - Pesquisa Assistida por Inteligência Artificial')
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rodape.runs[0].font.size = Pt(9)
    rodape.runs[0].font.color.rgb = RGBColor(128, 128, 128)
    
    # Salvar num buffer de memória
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def gerar_docx_intake(dados):
    """Gera um arquivo .docx formatado para Ficha de Intake"""
    doc = Document()
    
    titulo = doc.add_heading('Ficha de Intake (Resumo de Autos)', 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Adicionar os dados do JSON em formato limpo
    doc.add_heading('Partes', level=1)
    doc.add_paragraph(f"Autor: {dados.get('parte_autora', 'N/A')}")
    doc.add_paragraph(f"Réu: {dados.get('parte_re', 'N/A')}")
    
    doc.add_heading('Informações da Ação', level=1)
    doc.add_paragraph(f"Natureza da Ação: {dados.get('natureza_acao', 'N/A')}")
    doc.add_paragraph(f"Valor da Causa: {dados.get('valor_causa', 'N/A')}")
    
    doc.add_heading('Síntese dos Fatos', level=1)
    doc.add_paragraph(dados.get('sintese_fatos', 'N/A'))
    
    doc.add_heading('Pedidos Principais', level=1)
    for p in dados.get('pedidos_principais', []):
        doc.add_paragraph(f"• {p}")
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
